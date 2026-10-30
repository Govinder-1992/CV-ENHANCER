"""Function-by-function tests for the CV Enhancer backend.

Exercises the REAL functions (validation, text extraction, Bedrock JSON
parsing, HTTP responses, error mapping, worker batch handling) with AWS calls
mocked. Run: python tests/test_functions.py
"""
import base64
import io
import json
import os
import sys
from unittest import mock

# Make 'src' importable as the Lambda does (app.*, common.*).
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "src"))

# Env vars the modules read at import time.
os.environ.setdefault("RESULTS_BUCKET", "test-bucket")
os.environ.setdefault("NOTIFICATION_TOPIC_ARN", "arn:aws:sns:us-east-1:111111111111:test")
os.environ.setdefault("JOB_QUEUE_URL", "https://sqs.test/queue")
os.environ.setdefault("BEDROCK_MODEL_ID", "us.anthropic.claude-haiku-4-5-20251001-v1:0")

PASS, FAIL = 0, 0


def check(name, cond):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  PASS  {name}")
    else:
        FAIL += 1
        print(f"  FAIL  {name}")


def expect_error(name, fn, exc_substr):
    try:
        fn()
        check(name, False)
    except Exception as e:  # noqa: BLE001
        check(name, exc_substr.lower() in type(e).__name__.lower() or exc_substr.lower() in str(e).lower())


# ---------- build real sample files ----------
def make_pdf_bytes(text="John Doe\nSenior Engineer\nPython, AWS, Lambda"):
    from pypdf import PdfWriter
    import pypdf
    # pypdf can't easily author text; use reportlab-free approach: a minimal PDF
    # via a blank page won't contain text, so we instead craft with fpdf if present.
    try:
        from fpdf import FPDF
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("helvetica", size=12)
        for line in text.split("\n"):
            pdf.cell(0, 10, line, ln=1)
        return bytes(pdf.output())
    except Exception:
        return None  # signal: skip real-PDF extraction


def make_docx_bytes(text_lines=("Jane Roe", "Frontend Developer", "React, TypeScript, CSS")):
    import docx
    d = docx.Document()
    for line in text_lines:
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


PDF_MAGIC = b"%PDF-1.4\n"
ZIP_MAGIC = b"PK\x03\x04"


# ================= validation =================
def test_validation():
    print("\n[validation.py]")
    from common import validation
    from common.errors import ValidationError, UnsupportedFileTypeError, FileTooLargeError

    docx_bytes = make_docx_bytes()
    good = {
        "fileName": "cv.docx",
        "fileType": "docx",
        "fileContentBase64": base64.b64encode(docx_bytes).decode(),
        "jobRole": "Frontend Developer",
        "jobDescription": "We want a strong React developer with TypeScript experience.",
        "email": "jane@example.com",
    }
    data = validation.validate_payload(dict(good), 5 * 1024 * 1024)
    check("valid DOCX payload accepted", data["fileExt"] == "docx" and data["fileSize"] > 0)
    check("returns normalized fields", data["email"] == "jane@example.com" and data["jobRole"] == "Frontend Developer")

    expect_error("missing fileName rejected",
                 lambda: validation.validate_payload({**good, "fileName": ""}, 5_000_000), "validation")
    expect_error("bad email rejected",
                 lambda: validation.validate_payload({**good, "email": "not-an-email"}, 5_000_000), "validation")
    expect_error("short jobDescription rejected",
                 lambda: validation.validate_payload({**good, "jobDescription": "too short"}, 5_000_000), "validation")

    # Wrong magic bytes: claim pdf but send docx content
    expect_error("type/content mismatch rejected",
                 lambda: validation.validate_payload({**good, "fileType": "pdf"}, 5_000_000), "unsupported")

    # Too large
    big = base64.b64encode(ZIP_MAGIC + b"x" * (6 * 1024 * 1024)).decode()
    expect_error("oversized file rejected",
                 lambda: validation.validate_payload({**good, "fileContentBase64": big}, 5 * 1024 * 1024), "toolarge")

    # Unsupported extension
    expect_error("unsupported extension rejected",
                 lambda: validation.validate_payload({**good, "fileType": "txt", "fileName": "cv.txt"}, 5_000_000),
                 "unsupported")

    # Bad base64
    expect_error("invalid base64 rejected",
                 lambda: validation.validate_payload({**good, "fileContentBase64": "!!!notbase64!!!"}, 5_000_000),
                 "validation")


# ================= text extraction =================
def test_extraction():
    print("\n[text_extraction.py]")
    from common import text_extraction
    from common.errors import TextExtractionError

    # DOCX round-trip
    docx_bytes = make_docx_bytes(("Alice Smith", "DevOps Engineer", "Terraform, Kubernetes"))
    text = text_extraction.extract_text(docx_bytes, "docx")
    check("DOCX text extracted", "Alice Smith" in text and "Kubernetes" in text)

    # PDF round-trip (only if we could author a PDF)
    pdf_bytes = make_pdf_bytes("Bob Jones\nData Scientist\nPandas, SQL")
    if pdf_bytes:
        ptext = text_extraction.extract_text(pdf_bytes, "pdf")
        check("PDF text extracted", "Bob Jones" in ptext or "Data Scientist" in ptext)
    else:
        print("  SKIP  PDF extraction (no PDF authoring lib installed)")

    # Empty/garbage -> error
    expect_error("empty docx raises", lambda: text_extraction.extract_text(ZIP_MAGIC, "docx"), "extraction")
    expect_error("unsupported ext raises", lambda: text_extraction.extract_text(b"x", "rtf"), "extraction")


# ================= bedrock JSON parsing =================
def test_bedrock_parsing():
    print("\n[bedrock_client.py]")
    from app import bedrock_client
    from common.errors import BedrockError, TransientError
    from botocore.exceptions import ClientError

    model_json = {
        "matchScore": 82,
        "summary": "Strong fit.",
        "strengths": ["Python", "AWS"],
        "missingSkills": ["Kubernetes"],
        "recommendedKeywords": ["CI/CD", "microservices"],
        "gaps": ["No container orchestration"],
        "suggestions": ["Add a metrics bullet"],
        "enhancedCv": "John Doe\nSenior Engineer\n...",
    }

    # Mock a successful converse response wrapping the model's JSON (with fences).
    fake_ok = {
        "output": {"message": {"content": [{"text": "```json\n" + json.dumps(model_json) + "\n```"}]}},
        "usage": {"inputTokens": 100, "outputTokens": 200},
    }
    with mock.patch.object(bedrock_client._bedrock, "converse", return_value=fake_ok):
        out = bedrock_client.enhance_cv("cv text", "Engineer", "job desc long enough")
    check("enhance_cv parses analysis", out["analysis"]["matchScore"] == 82)
    check("enhance_cv returns new fields",
          out["analysis"]["missingSkills"] == ["Kubernetes"] and out["analysis"]["recommendedKeywords"] == ["CI/CD", "microservices"])
    check("enhance_cv returns enhancedCv", out["enhancedCv"].startswith("John Doe"))

    # Throttling -> TransientError
    throttle = ClientError({"Error": {"Code": "ThrottlingException", "Message": "slow down"}}, "Converse")
    with mock.patch.object(bedrock_client._bedrock, "converse", side_effect=throttle):
        expect_error("throttling -> TransientError",
                     lambda: bedrock_client.enhance_cv("t", "r", "job desc long enough"), "transient")

    # Non-JSON model output -> BedrockError
    fake_bad = {"output": {"message": {"content": [{"text": "sorry, no json here"}]}}, "usage": {}}
    with mock.patch.object(bedrock_client._bedrock, "converse", return_value=fake_bad):
        expect_error("bad model output -> BedrockError",
                     lambda: bedrock_client.enhance_cv("t", "r", "job desc long enough"), "bedrock")


# ================= http responses =================
def test_http():
    print("\n[http.py]")
    from common import http
    r = http.response(200, {"ok": True})
    body = json.loads(r["body"])
    check("response wraps status + body", r["statusCode"] == 200 and body["ok"] is True)
    check("CORS header present", "Access-Control-Allow-Origin" in r["headers"])
    e = http.error(400, "validation_error", "bad input", requestId="abc")
    eb = json.loads(e["body"])
    check("error includes code+message+extra", eb["error"] == "validation_error" and eb["requestId"] == "abc")


# ================= errors mapping =================
def test_errors():
    print("\n[errors.py]")
    from common import errors
    check("ValidationError -> 400", errors.ValidationError().status_code == 400)
    check("FileTooLargeError code", errors.FileTooLargeError().code == "file_too_large")
    check("TextExtractionError -> 422", errors.TextExtractionError().status_code == 422)
    check("TransientError is CvEnhancerError", issubclass(errors.TransientError, errors.CvEnhancerError))


# ================= storage (mocked S3) =================
def test_storage():
    print("\n[storage.py]")
    from app import storage
    with mock.patch.object(storage._s3, "put_object", return_value={}) as put, \
         mock.patch.object(storage._s3, "generate_presigned_url", return_value="https://signed.example/cv.txt") as sign:
        res = storage.save_results("req-123", "ENHANCED CV TEXT", {"matchScore": 70}, {"email": "a@b.com"})
    check("save_results writes 2 objects (cv + analysis)", put.call_count == 2)
    check("save_results returns download url", res["downloadUrl"].startswith("https://"))
    check("save_results uses results/ prefix", res["cvKey"].startswith("results/req-123/"))
    check("put_object uses AES256 SSE",
          all(c.kwargs.get("ServerSideEncryption") == "AES256" for c in put.call_args_list))


# ================= notifications (mocked SNS) =================
def test_notifications():
    print("\n[notifications.py]")
    from app import notifications
    # list_subscriptions_by_topic paginator + subscribe + publish all mocked.
    fake_paginator = mock.Mock()
    fake_paginator.paginate.return_value = [{"Subscriptions": []}]
    with mock.patch.object(notifications._sns, "get_paginator", return_value=fake_paginator), \
         mock.patch.object(notifications._sns, "subscribe", return_value={"SubscriptionArn": "pending"}) as sub, \
         mock.patch.object(notifications._sns, "publish", return_value={"MessageId": "m1"}) as pub:
        notifications.notify_success("req-1", "user@example.com", "Engineer",
                                     {"downloadUrl": "https://x", "expiresInSeconds": 86400}, 75)
    check("notify_success subscribes new email", sub.called)
    check("notify_success publishes", pub.called)
    # The published message must NOT contain CV content, only metadata + link.
    published = json.loads(pub.call_args.kwargs["Message"])
    check("notification has download link", published.get("downloadUrl") == "https://x")
    check("notification carries matchScore", published.get("matchScore") == 75)
    check("notification body has no raw CV text", "cvText" not in published and "enhancedCv" not in published)


# ================= worker batch handling (mocked) =================
def test_worker():
    print("\n[worker.py]")
    from app import worker

    docx_b64 = base64.b64encode(make_docx_bytes()).decode()
    body = {
        "fileName": "cv.docx", "fileType": "docx", "fileContentBase64": docx_b64,
        "jobRole": "Frontend Developer",
        "jobDescription": "Looking for a React developer with strong TypeScript skills.",
        "email": "jane@example.com",
    }
    event = {"Records": [{"messageId": "mid-1", "body": json.dumps(body), "messageAttributes": {}}]}

    fake_enhanced = {"analysis": {"matchScore": 80}, "enhancedCv": "ENHANCED"}
    with mock.patch.object(worker.bedrock_client, "enhance_cv", return_value=fake_enhanced), \
         mock.patch.object(worker.storage, "save_results",
                           return_value={"downloadUrl": "https://x", "expiresInSeconds": 86400, "cvKey": "results/mid-1/enhanced-cv.txt"}), \
         mock.patch.object(worker.notifications, "notify_success", return_value=None):
        out = worker.lambda_handler(event, None)
    check("worker success -> no batch failures", out == {"batchItemFailures": []})

    # Transient error -> message reported for retry
    from common.errors import TransientError
    with mock.patch.object(worker.bedrock_client, "enhance_cv", side_effect=TransientError("throttle")), \
         mock.patch.object(worker.notifications, "notify_failure", return_value=None):
        out2 = worker.lambda_handler(event, None)
    check("worker transient -> reports batch failure", out2["batchItemFailures"] == [{"itemIdentifier": "mid-1"}])

    # Permanent (validation) error -> swallowed, no retry
    bad_event = {"Records": [{"messageId": "mid-2", "body": json.dumps({**body, "email": "bad"}), "messageAttributes": {}}]}
    with mock.patch.object(worker.notifications, "notify_failure", return_value=None):
        out3 = worker.lambda_handler(bad_event, None)
    check("worker permanent error -> not retried", out3 == {"batchItemFailures": []})


if __name__ == "__main__":
    for fn in (test_validation, test_extraction, test_bedrock_parsing, test_http,
               test_errors, test_storage, test_notifications, test_worker):
        try:
            fn()
        except Exception as e:  # noqa: BLE001
            FAIL += 1
            print(f"  ERROR in {fn.__name__}: {type(e).__name__}: {e}")

    print(f"\n==== RESULT: {PASS} passed, {FAIL} failed ====")
    sys.exit(1 if FAIL else 0)
